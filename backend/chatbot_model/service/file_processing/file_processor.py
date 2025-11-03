"""
Advanced File Processor for AI Chatbot Builder
Handles various file formats and data processing
"""

import json
import csv
import os
import logging
from typing import Dict, List, Any, Optional, Union
from pathlib import Path
import pandas as pd
import docx
import openpyxl
from io import StringIO

logger = logging.getLogger(__name__)

class AdvancedFileProcessor:
    """Advanced file processing with support for multiple formats"""
    
    def __init__(self):
        self.supported_formats = {
            '.json': self._process_json,
            '.csv': self._process_csv,
            '.txt': self._process_txt,
            '.xlsx': self._process_excel,
            '.xls': self._process_excel,
            '.docx': self._process_docx
        }
    
    def process_file(self, file_path: str, filename: str = None) -> Dict[str, Any]:
        """Process a file and extract intent data"""
        try:
            if filename is None:
                filename = os.path.basename(file_path)
            
            file_extension = Path(filename).suffix.lower()
            
            if file_extension not in self.supported_formats:
                raise ValueError(f"Unsupported file format: {file_extension}")
            
            processor = self.supported_formats[file_extension]
            result = processor(file_path)
            
            return {
                'success': True,
                'data': result,
                'format': file_extension,
                'filename': filename,
                'processed_count': len(result) if isinstance(result, list) else 1
            }
            
        except Exception as e:
            logger.error(f"Error processing file {filename}: {e}")
            return {
                'success': False,
                'error': str(e),
                'format': file_extension if 'file_extension' in locals() else 'unknown',
                'filename': filename or 'unknown'
            }
    
    def _process_json(self, file_path: str) -> List[Dict[str, Any]]:
        """Process JSON file containing intents"""
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        
        if isinstance(data, dict) and 'intents' in data:
            return data['intents']
        elif isinstance(data, list):
            return data
        else:
            # Try to convert single object to intent format
            if isinstance(data, dict) and 'tag' in data:
                return [data]
            else:
                raise ValueError("Invalid JSON format. Expected intents array or single intent object.")
    
    def _process_csv(self, file_path: str) -> List[Dict[str, Any]]:
        """Process CSV file with columns: tag, patterns, responses"""
        intents = []
        
        try:
            # Try pandas first for better CSV handling
            df = pd.read_csv(file_path, encoding='utf-8')
            
            # Check required columns
            required_cols = ['tag', 'patterns', 'responses']
            if not all(col in df.columns for col in required_cols):
                # Try alternative column names
                alt_mapping = {
                    'intent': 'tag',
                    'pattern': 'patterns',
                    'response': 'responses',
                    'question': 'patterns',
                    'answer': 'responses'
                }
                
                for alt_col, std_col in alt_mapping.items():
                    if alt_col in df.columns and std_col not in df.columns:
                        df = df.rename(columns={alt_col: std_col})
            
            # Process each row
            for _, row in df.iterrows():
                if pd.isna(row.get('tag')) or pd.isna(row.get('patterns')):
                    continue
                
                intent = {
                    'tag': str(row['tag']).strip(),
                    'patterns': self._split_field(str(row['patterns'])),
                    'responses': self._split_field(str(row.get('responses', '')))
                }
                
                if intent['patterns'] and intent['responses']:
                    intents.append(intent)
        
        except Exception as e:
            logger.warning(f"Pandas CSV processing failed: {e}. Trying basic CSV reader.")
            
            # Fallback to basic CSV reader
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.DictReader(f)
                
                for row in reader:
                    if 'tag' in row and 'patterns' in row:
                        intent = {
                            'tag': row['tag'].strip(),
                            'patterns': self._split_field(row['patterns']),
                            'responses': self._split_field(row.get('responses', ''))
                        }
                        
                        if intent['patterns'] and intent['responses']:
                            intents.append(intent)
        
        return intents
    
    def _process_txt(self, file_path: str) -> List[Dict[str, Any]]:
        """Process text file with structured format"""
        intents = []
        current_intent = None
        
        with open(file_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        
        for line in lines:
            line = line.strip()
            if not line or line.startswith('#'):  # Skip empty lines and comments
                continue
            
            if line.startswith('TAG:') or line.startswith('INTENT:'):
                if current_intent:
                    intents.append(current_intent)
                current_intent = {
                    'tag': line.split(':', 1)[1].strip(),
                    'patterns': [],
                    'responses': []
                }
            elif line.startswith('PATTERN:') or line.startswith('Q:'):
                if current_intent:
                    pattern = line.split(':', 1)[1].strip()
                    if pattern:
                        current_intent['patterns'].append(pattern)
            elif line.startswith('RESPONSE:') or line.startswith('A:'):
                if current_intent:
                    response = line.split(':', 1)[1].strip()
                    if response:
                        current_intent['responses'].append(response)
            elif line.startswith('PATTERNS:'):
                if current_intent:
                    patterns_text = line.split(':', 1)[1].strip()
                    current_intent['patterns'].extend(self._split_field(patterns_text))
            elif line.startswith('RESPONSES:'):
                if current_intent:
                    responses_text = line.split(':', 1)[1].strip()
                    current_intent['responses'].extend(self._split_field(responses_text))
        
        if current_intent:
            intents.append(current_intent)
        
        return intents
    
    def _process_excel(self, file_path: str) -> List[Dict[str, Any]]:
        """Process Excel file with columns: tag, patterns, responses"""
        intents = []
        
        try:
            # Try pandas first
            df = pd.read_excel(file_path)
            
            # Check and map columns
            required_cols = ['tag', 'patterns', 'responses']
            if not all(col in df.columns for col in required_cols):
                # Try alternative column names
                alt_mapping = {
                    'intent': 'tag',
                    'pattern': 'patterns',
                    'response': 'responses',
                    'question': 'patterns',
                    'answer': 'responses'
                }
                
                for alt_col, std_col in alt_mapping.items():
                    if alt_col in df.columns and std_col not in df.columns:
                        df = df.rename(columns={alt_col: std_col})
            
            # Process each row
            for _, row in df.iterrows():
                if pd.isna(row.get('tag')) or pd.isna(row.get('patterns')):
                    continue
                
                intent = {
                    'tag': str(row['tag']).strip(),
                    'patterns': self._split_field(str(row['patterns'])),
                    'responses': self._split_field(str(row.get('responses', '')))
                }
                
                if intent['patterns'] and intent['responses']:
                    intents.append(intent)
        
        except Exception as e:
            logger.warning(f"Pandas Excel processing failed: {e}. Trying openpyxl.")
            
            # Fallback to openpyxl
            try:
                workbook = openpyxl.load_workbook(file_path)
                sheet = workbook.active
                
                # Get header row
                headers = [cell.value for cell in sheet[1] if cell.value]
                
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if len(row) >= 3 and row[0]:
                        intent = {
                            'tag': str(row[0]).strip(),
                            'patterns': self._split_field(str(row[1] or '')),
                            'responses': self._split_field(str(row[2] or ''))
                        }
                        
                        if intent['patterns'] and intent['responses']:
                            intents.append(intent)
            
            except Exception as e2:
                logger.error(f"Excel processing failed completely: {e2}")
                raise
        
        return intents
    
    def _process_docx(self, file_path: str) -> List[Dict[str, Any]]:
        """Process Word document with structured content"""
        intents = []
        current_intent = None
        
        try:
            doc = docx.Document(file_path)
            
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                
                if text.startswith('TAG:') or text.startswith('INTENT:'):
                    if current_intent:
                        intents.append(current_intent)
                    current_intent = {
                        'tag': text.split(':', 1)[1].strip(),
                        'patterns': [],
                        'responses': []
                    }
                elif text.startswith('PATTERN:') or text.startswith('Q:'):
                    if current_intent:
                        pattern = text.split(':', 1)[1].strip()
                        if pattern:
                            current_intent['patterns'].append(pattern)
                elif text.startswith('RESPONSE:') or text.startswith('A:'):
                    if current_intent:
                        response = text.split(':', 1)[1].strip()
                        if response:
                            current_intent['responses'].append(response)
                elif text.startswith('PATTERNS:'):
                    if current_intent:
                        patterns_text = text.split(':', 1)[1].strip()
                        current_intent['patterns'].extend(self._split_field(patterns_text))
                elif text.startswith('RESPONSES:'):
                    if current_intent:
                        responses_text = text.split(':', 1)[1].strip()
                        current_intent['responses'].extend(self._split_field(responses_text))
            
            if current_intent:
                intents.append(current_intent)
        
        except Exception as e:
            logger.error(f"Error processing Word document: {e}")
            raise
        
        return intents
    
    def _split_field(self, field_text: str) -> List[str]:
        """Split field text by various delimiters"""
        if not field_text or field_text == 'nan':
            return []
        
        # Try different delimiters
        delimiters = ['|', ';', '\n', '\\n']
        
        for delimiter in delimiters:
            if delimiter in field_text:
                parts = [part.strip() for part in field_text.split(delimiter)]
                return [part for part in parts if part]
        
        # If no delimiter found, return as single item
        return [field_text.strip()] if field_text.strip() else []
    
    def validate_intents(self, intents: List[Dict[str, Any]]) -> Dict[str, Any]:
        """Validate processed intents data"""
        validation_result = {
            'valid': True,
            'total_intents': len(intents),
            'valid_intents': 0,
            'issues': [],
            'warnings': []
        }
        
        valid_intents = []
        
        for i, intent in enumerate(intents):
            intent_issues = []
            
            # Check required fields
            if not intent.get('tag'):
                intent_issues.append(f"Intent {i+1}: Missing tag")
            
            if not intent.get('patterns'):
                intent_issues.append(f"Intent {i+1}: Missing patterns")
            elif not isinstance(intent['patterns'], list):
                intent_issues.append(f"Intent {i+1}: Patterns should be a list")
            
            if not intent.get('responses'):
                intent_issues.append(f"Intent {i+1}: Missing responses")
            elif not isinstance(intent['responses'], list):
                intent_issues.append(f"Intent {i+1}: Responses should be a list")
            
            # Check for empty patterns/responses
            if intent.get('patterns'):
                empty_patterns = [j for j, p in enumerate(intent['patterns']) if not p.strip()]
                if empty_patterns:
                    validation_result['warnings'].append(
                        f"Intent {i+1}: Empty patterns at positions {empty_patterns}"
                    )
            
            if intent.get('responses'):
                empty_responses = [j for j, r in enumerate(intent['responses']) if not r.strip()]
                if empty_responses:
                    validation_result['warnings'].append(
                        f"Intent {i+1}: Empty responses at positions {empty_responses}"
                    )
            
            if intent_issues:
                validation_result['issues'].extend(intent_issues)
                validation_result['valid'] = False
            else:
                valid_intents.append(intent)
                validation_result['valid_intents'] += 1
        
        validation_result['processed_intents'] = valid_intents
        
        return validation_result
    
    def export_intents(self, intents: List[Dict[str, Any]], output_path: str, format: str = 'json') -> bool:
        """Export intents to various formats"""
        try:
            if format.lower() == 'json':
                with open(output_path, 'w', encoding='utf-8') as f:
                    json.dump({'intents': intents}, f, indent=2, ensure_ascii=False)
            
            elif format.lower() == 'csv':
                with open(output_path, 'w', newline='', encoding='utf-8') as f:
                    writer = csv.writer(f)
                    writer.writerow(['tag', 'patterns', 'responses'])
                    
                    for intent in intents:
                        writer.writerow([
                            intent.get('tag', ''),
                            '|'.join(intent.get('patterns', [])),
                            '|'.join(intent.get('responses', []))
                        ])
            
            elif format.lower() == 'txt':
                with open(output_path, 'w', encoding='utf-8') as f:
                    for intent in intents:
                        f.write(f"TAG: {intent.get('tag', '')}\n")
                        for pattern in intent.get('patterns', []):
                            f.write(f"PATTERN: {pattern}\n")
                        for response in intent.get('responses', []):
                            f.write(f"RESPONSE: {response}\n")
                        f.write("\n")
            
            else:
                raise ValueError(f"Unsupported export format: {format}")
            
            return True
            
        except Exception as e:
            logger.error(f"Error exporting intents: {e}")
            return False

# Global instance
file_processor = AdvancedFileProcessor()

# Convenience functions
def process_file(file_path: str, filename: str = None) -> Dict[str, Any]:
    return file_processor.process_file(file_path, filename)

def validate_intents(intents: List[Dict[str, Any]]) -> Dict[str, Any]:
    return file_processor.validate_intents(intents)

def export_intents(intents: List[Dict[str, Any]], output_path: str, format: str = 'json') -> bool:
    return file_processor.export_intents(intents, output_path, format)