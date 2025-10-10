
import os
import json
import csv
from pathlib import Path
import openpyxl
import docx
  
from config import logger

def process_multiple_file_formats(file_path: str, filename: str):
    """Process various file formats and extract intent data"""
    try:
        file_extension = Path(filename).suffix.lower()
        
        if file_extension == '.json':
            return process_json_file(file_path)
        elif file_extension == '.csv':
            return process_csv_file(file_path)
        elif file_extension == '.txt':
            return process_txt_file(file_path)
        elif file_extension == '.xlsx' or file_extension == '.xls':
            return process_excel_file(file_path)
        elif file_extension == '.docx':
            return process_docx_file(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")
            
    except Exception as e:
        logger.error(f"Error processing file {filename}: {e}")
        raise

def process_json_file(file_path: str):
    """Process JSON file containing intents"""
    with open(file_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    if 'intents' in data:
        return data['intents']
    elif isinstance(data, list):
        return data
    else:
        # Try to convert single object to intent format
        return [data] if 'tag' in data else []

def process_csv_file(file_path: str):
    """Process CSV file with columns: tag, patterns, responses"""
    intents = []
    
    with open(file_path, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        
        for row in reader:
            if 'tag' in row and 'patterns' in row and 'responses' in row:
                intent = {
                    'tag': row['tag'],
                    'patterns': [p.strip() for p in row['patterns'].split('|') if p.strip()],
                    'responses': [r.strip() for r in row['responses'].split('|') if r.strip()]
                }
                intents.append(intent)
    
    return intents

def process_txt_file(file_path: str):
    """Process text file with simple format"""
    intents = []
    current_intent = None
    
    with open(file_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('TAG:'):
            if current_intent:
                intents.append(current_intent)
            current_intent = {
                'tag': line[4:].strip(),
                'patterns': [],
                'responses': []
            }
        elif line.startswith('PATTERN:') and current_intent:
            current_intent['patterns'].append(line[8:].strip())
        elif line.startswith('RESPONSE:') and current_intent:
            current_intent['responses'].append(line[9:].strip())
    
    if current_intent:
        intents.append(current_intent)
    
    return intents

def process_excel_file(file_path: str):
    """Process Excel file with columns: tag, patterns, responses"""
    intents = []
    
    try:
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        
        # Get header row
        headers = [cell.value for cell in sheet[1]]
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) >= 3 and row[0]:  # Ensure we have at least 3 columns and tag is not empty
                intent = {
                    'tag': str(row[0]).strip(),
                    'patterns': [p.strip() for p in str(row[1] or '').split('|') if p.strip()],
                    'responses': [r.strip() for r in str(row[2] or '').split('|') if r.strip()]
                }
                if intent['patterns'] and intent['responses']:
                    intents.append(intent)
                    
    except Exception as e:
        logger.error(f"Error processing Excel file: {e}")
        raise
    
    return intents

def process_docx_file(file_path: str):
    """Process Word document with structured content"""
    intents = []
    
    try:
        doc = docx.Document(file_path)
        current_intent = None
        
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
                
            if text.startswith('TAG:'):
                if current_intent:
                    intents.append(current_intent)
                current_intent = {
                    'tag': text[4:].strip(),
                    'patterns': [],
                    'responses': []
                }
            elif text.startswith('PATTERN:') and current_intent:
                current_intent['patterns'].append(text[8:].strip())
            elif text.startswith('RESPONSE:') and current_intent:
                current_intent['responses'].append(text[9:].strip())
        
        if current_intent:
            intents.append(current_intent)
            
    except Exception as e:
        logger.error(f"Error processing Word document: {e}")
        raise
    
    return intents
